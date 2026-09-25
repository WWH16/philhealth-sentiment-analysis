import os
import json
from datetime import datetime, time, timedelta
from collections import defaultdict
from functools import wraps


from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib.auth.models import User, Group, Permission
from django.contrib.auth.hashers import make_password
from django.contrib.admin.models import LogEntry, ADDITION, CHANGE, DELETION
from django.contrib.contenttypes.models import ContentType
from django.db import transaction
from django.db.models import Count, Q
from django.contrib import messages
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import csrf_exempt
from django.conf import settings
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.utils import timezone
from django.db.models.functions import TruncDate, TruncWeek, TruncMonth, ExtractHour
from feedback.models import FeedbackConfiguration, FeedbackEntry
from feedback.email_service import send_daily_summary_email

from django.http import FileResponse, Http404
from django.core.exceptions import SuspiciousFileOperation

from feedback_admin.backup_utils import (
    create_backup, list_backups, resolve_backup_path, delete_backup, restore_backup,
)

# ── Activity log helpers (built on Django's built-in django_admin_log table) ──

def _feedback_content_type():
    return ContentType.objects.get_for_model(FeedbackEntry)


def log_feedback_note(user, entry, body):
    """Records an 'action taken / reply' note as an ADDITION log entry."""
    LogEntry.objects.log_actions(
        user_id=user.id,
        queryset=[entry],
        action_flag=ADDITION,
        change_message=body,
        single_object=True,
    )


def log_feedback_status_change(user, entry, old_status, new_status):
    """Records a status change as a CHANGE log entry."""
    LogEntry.objects.log_actions(
        user_id=user.id,
        queryset=[entry],
        action_flag=CHANGE,
        change_message=f'{old_status}|{new_status}',
        single_object=True,
    )


def log_admin_event(user, obj, action_flag, change_message):
    """Writes a generic built-in admin log entry for any tracked object."""
    LogEntry.objects.log_actions(
        user_id=user.id,
        queryset=[obj],
        action_flag=action_flag,
        change_message=change_message,
        single_object=True,
    )


def get_feedback_activity(entry):
    """Builds (notes, status_history) for an entry from django_admin_log."""
    logs = (LogEntry.objects
            .filter(content_type=_feedback_content_type(), object_id=str(entry.pk))
            .select_related('user')
            .order_by('action_time'))

    status_display = dict(FeedbackEntry.STATUS_CHOICES)
    notes, history = [], []

    for log in logs:
        author = (log.user.get_full_name() or log.user.username) if log.user else 'System'
        at = timezone.localtime(log.action_time).strftime('%b %d, %Y %I:%M %p')

        if log.action_flag == CHANGE and '|' in log.change_message:
            old_raw, _, new_raw = log.change_message.partition('|')
            history.append({
                'old': status_display.get(old_raw, old_raw),
                'new': status_display.get(new_raw, new_raw),
                'by': author,
                'at': at,
            })
        else:
            notes.append({'author': author, 'body': log.change_message, 'created_at': at})

    return notes, history


def _build_feedback_activity_map(entry_ids):
    """Builds feedback activity for many entries in one query."""
    if not entry_ids:
        return {}

    logs = (LogEntry.objects
            .filter(content_type=_feedback_content_type(), object_id__in=[str(pk) for pk in entry_ids])
            .select_related('user')
            .order_by('object_id', 'action_time'))

    status_display = dict(FeedbackEntry.STATUS_CHOICES)
    activity = {}

    for log in logs:
        if not log.object_id or not log.object_id.isdigit():
            continue

        entry_id = int(log.object_id)
        bucket = activity.setdefault(entry_id, {'notes': [], 'status_history': []})
        author = (log.user.get_full_name() or log.user.username) if log.user else 'System'
        at = timezone.localtime(log.action_time).strftime('%b %d, %Y %I:%M %p')

        if log.action_flag == CHANGE and '|' in log.change_message:
            old_raw, _, new_raw = log.change_message.partition('|')
            bucket['status_history'].append({
                'old': status_display.get(old_raw, old_raw),
                'new': status_display.get(new_raw, new_raw),
                'by': author,
                'at': at,
            })
        else:
            bucket['notes'].append({
                'author': author,
                'body': log.change_message,
                'created_at': at,
            })

    return activity


def _format_audit_row(log, entry=None):
    local_time = timezone.localtime(log.action_time)
    username = log.user.username if log.user else 'System'
    full_name = log.user.get_full_name() if log.user else ''
    ctype = log.content_type.model if log.content_type_id else ''
    message = log.change_message or ''
    action_label = 'Event'
    summary = message
    action_type = 'event'
    search_parts = [
        username,
        full_name,
        message,
        ctype,
        log.object_repr or '',
        f'Feedback #{entry.pk}' if entry else '',
        getattr(entry, 'username', ''),
        getattr(entry, 'name', ''),
    ]
    old_status = ''
    new_status = ''

    if ctype == 'feedbackentry':
        label = f"Feedback #{entry.pk}" if entry else log.object_repr
        if log.action_flag == CHANGE and '|' in message:
            old_raw, _, new_raw = message.partition('|')
            action_label = 'Status Update'
            action_type = 'status'
            old_status = dict(FeedbackEntry.STATUS_CHOICES).get(old_raw, old_raw)
            new_status = dict(FeedbackEntry.STATUS_CHOICES).get(new_raw, new_raw)
            summary = f"{label}: {old_status} -> {new_status}"
        elif log.action_flag == CHANGE and message.startswith('category:'):
            _, _, payload = message.partition(':')
            old_raw, _, new_raw = payload.partition('|')
            action_label = 'Category Update'
            action_type = 'category'
            summary = f"{label}: {dict(FeedbackEntry.CATEGORY_CHOICES).get(old_raw, old_raw or 'Uncategorized')} -> {dict(FeedbackEntry.CATEGORY_CHOICES).get(new_raw, new_raw or 'Uncategorized')}"
        elif log.action_flag == DELETION:
            action_label = 'Feedback Deleted'
            action_type = 'delete'
            summary = f"{label} deleted"
        elif log.action_flag == ADDITION:
            action_label = 'Note / Reply'
            action_type = 'note'
            summary = f"{label}: {message}"
        else:
            summary = f"{label}: {message}"
    elif ctype == 'user':
        if log.action_flag == ADDITION and message == 'Logged in':
            action_label = 'Login'
            action_type = 'login'
            summary = f"{log.object_repr} logged in"
        elif log.action_flag == CHANGE and message == 'Logged out':
            action_label = 'Logout'
            action_type = 'logout'
            summary = f"{log.object_repr} logged out"
        elif log.action_flag == ADDITION:
            action_label = 'User Created'
            action_type = 'create'
            summary = message or f'{log.object_repr} created'
        elif log.action_flag == CHANGE and message == 'Password updated':
            action_label = 'Password Updated'
            action_type = 'profile'
            summary = f"{log.object_repr}: password updated"
        elif log.action_flag == CHANGE:
            action_label = 'Profile Updated'
            action_type = 'profile'
            summary = message or f'{log.object_repr} updated'
        elif log.action_flag == DELETION:
            action_label = 'User Deleted'
            action_type = 'delete'
            summary = message or f'{log.object_repr} deleted'
    elif ctype == 'group':
        if log.action_flag == ADDITION:
            action_label = 'Group Created'
            action_type = 'create'
        elif log.action_flag == CHANGE:
            action_label = 'Group Updated'
            action_type = 'update'
        elif log.action_flag == DELETION:
            action_label = 'Group Deleted'
            action_type = 'delete'
        summary = message or log.object_repr
    elif ctype == 'feedbackconfiguration':
        if message.startswith('Auto-analysis'):
            action_label = 'Auto-Analysis Toggle'
            action_type = 'settings'
        elif message.startswith('Batch re-analysis'):
            action_label = 'Batch Re-analyze'
            action_type = 'settings'
        elif message.startswith('Created backup'):
            action_label = 'Backup Created'
            action_type = 'backup'
        elif message.startswith('Deleted backup'):
            action_label = 'Backup Deleted'
            action_type = 'backup'
        elif message.startswith('Restored database'):
            action_label = 'Database Restored'
            action_type = 'backup'
        else:
            action_label = 'Settings Update'
            action_type = 'settings'
        summary = message

    return {
        'id': log.pk,
        'date': local_time.strftime('%Y-%m-%d'),
        'time': local_time.strftime('%H:%M'),
        'admin': username,
        'username': username,
        'full_name': full_name,
        'action_type': action_type,
        'action_label': action_label,
        'summary': summary,
        'old_status': old_status,
        'new_status': new_status,
        'search_text': ' '.join(str(part) for part in search_parts if part),
    }


def _is_ajax(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest'


def staff_required(view_func):
    """Requires user to be authenticated and have staff or superuser status."""
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('admin_login')
        if not (request.user.is_staff or request.user.is_superuser):
            messages.error(request, 'Access denied. Staff privileges are required.')
            return redirect('admin_login')
        return view_func(request, *args, **kwargs)
    return _wrapped_view


def superuser_required(view_func):
    """Requires user to be authenticated and have superuser status."""
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('admin_login')
        if not request.user.is_superuser:
            if _is_ajax(request):
                return JsonResponse({'ok': False, 'error': 'Access restricted to administrators.'}, status=403)
            messages.error(request, 'Access restricted to administrators.')
            return redirect('dashboard')
        return view_func(request, *args, **kwargs)
    return _wrapped_view


@staff_required
@require_POST
def response_note_add(request, entry_id):
    entry = get_object_or_404(FeedbackEntry, pk=entry_id)
    try:
        payload = json.loads(request.body.decode('utf-8'))
    except json.JSONDecodeError:
        return JsonResponse({'ok': False, 'error': 'Invalid request.'}, status=400)

    body = (payload.get('body') or '').strip()
    if not body:
        return JsonResponse({'ok': False, 'error': 'Note body cannot be empty.'}, status=400)

    log_feedback_note(request.user, entry, body)

    return JsonResponse({
        'ok': True,
        'note': {
            'author': request.user.get_full_name() or request.user.username,
            'body': body,
            'created_at': timezone.localtime(timezone.now()).strftime('%b %d, %Y %I:%M %p'),
        }
    })


@staff_required
def dashboard(request):
    qs = FeedbackEntry.objects.all()
    now = timezone.localtime(timezone.now())
    today = now.date()
    week_start = now - timedelta(days=7)
    month_start = now - timedelta(days=30)
    recent_entries = qs.select_related('staff_assisted').order_by('-created_at')[:5]

    today_start = timezone.make_aware(datetime.combine(today, time.min))
    today_end = timezone.make_aware(datetime.combine(today, time.max))

    filter_data = _multi_period_experience_counts(qs, (today_start, today_end), week_start, month_start)
    all_counts = filter_data['all']
    total = all_counts['total']
    sa = all_counts['strongly_agree']
    a = all_counts['agree']
    nad = all_counts['neither']
    d = all_counts['disagree']
    sd = all_counts['strongly_disagree']
    na = all_counts['na']

    trend_counts = (
        qs.annotate(local_date=TruncDate('created_at'))
        .values('local_date', 'experience')
        .annotate(count=Count('id'))
    )
    trend_data_map = defaultdict(lambda: defaultdict(int))
    for row in trend_counts:
        if row['local_date']:
            exp = row['experience']
            if exp == 'vsat':
                exp = FeedbackEntry.STRONGLY_AGREE
            elif exp == 'sat':
                exp = FeedbackEntry.AGREE
            elif exp == 'unsat':
                exp = FeedbackEntry.STRONGLY_DISAGREE
            trend_data_map[row['local_date']][exp] += row['count']

    if trend_data_map:
        earliest_date = min(trend_data_map.keys())
        start_date = min(earliest_date, today - timedelta(days=6))
        total_days = (today - start_date).days + 1
        trend_dates = [start_date + timedelta(days=i) for i in range(total_days)]
    else:
        trend_dates = [today - timedelta(days=i) for i in range(6, -1, -1)]

    def pct(value):
        return round((value / total) * 100) if total else 0

    def trend_for(experience):
        return [trend_data_map[day][experience] for day in trend_dates]

    recent_entries_data = list(recent_entries)
    recent_activity = _build_feedback_activity_map([entry.pk for entry in recent_entries_data])

    context = {
        'total': total,
        'strongly_agree': sa,
        'agree': a,
        'neither': nad,
        'disagree': d,
        'strongly_disagree': sd,
        'na': na,
        'strongly_agree_pct': pct(sa),
        'agree_pct': pct(a),
        'neither_pct': pct(nad),
        'disagree_pct': pct(d),
        'strongly_disagree_pct': pct(sd),
        'na_pct': pct(na),
        'recent_entries_data': [_entry_to_row(entry, recent_activity) for entry in recent_entries_data],
        'filter_data': filter_data,
        'trend_labels': [f'{day:%b} {day.day}' for day in trend_dates],
        'trend_dates': [day.isoformat() for day in trend_dates],
        'trend_strongly_agree': trend_for(FeedbackEntry.STRONGLY_AGREE),
        'trend_agree': trend_for(FeedbackEntry.AGREE),
        'trend_neither': trend_for(FeedbackEntry.NEITHER),
        'trend_disagree': trend_for(FeedbackEntry.DISAGREE),
        'trend_strongly_disagree': trend_for(FeedbackEntry.STRONGLY_DISAGREE),
        'trend_na': trend_for(FeedbackEntry.NOT_APPLICABLE),
    }
    return render(request, 'feedback_admin/dashboard.html', context)


@staff_required
def responses(request):
    entries = FeedbackEntry.objects.select_related('staff_assisted').order_by('-created_at')
    counts = _experience_counts(entries)
    entries_data = list(entries)
    activity_map = _build_feedback_activity_map([entry.pk for entry in entries_data])

    context = {
        'total': counts['total'],
        'strongly_agree': counts['strongly_agree'],
        'agree': counts['agree'],
        'neither': counts['neither'],
        'disagree': counts['disagree'],
        'strongly_disagree': counts['strongly_disagree'],
        'na': counts['na'],
        'entries_data': [_entry_to_row(entry, activity_map) for entry in entries_data],
    }
    return render(request, 'feedback_admin/responses.html', context)


_EXP_DISPLAY = dict(FeedbackEntry.EXPERIENCE_CHOICES)
_CAT_DISPLAY = dict(FeedbackEntry.CATEGORY_CHOICES)
_STATUS_DISPLAY = dict(FeedbackEntry.STATUS_CHOICES)
_SENT_DISPLAY = dict(FeedbackEntry.SENTIMENT_CHOICES)


def _entry_to_row(entry, activity=None):
    local_created = timezone.localtime(entry.created_at)
    if activity is None:
        notes, status_history = get_feedback_activity(entry)
    else:
        entry_activity = activity.get(entry.pk, {'notes': [], 'status_history': []})
        notes = entry_activity['notes']
        status_history = entry_activity['status_history']
    has_comment = bool(entry.comment and entry.comment.strip())
    if not has_comment or entry.sentiment == FeedbackEntry.NOT_APPLICABLE:
        sentiment_display = 'N/A'
        sentiment_value = FeedbackEntry.NOT_APPLICABLE
    else:
        sentiment_display = _SENT_DISPLAY.get(entry.sentiment, entry.sentiment)
        sentiment_value = entry.sentiment

    exp_display = _EXP_DISPLAY.get(entry.experience, entry.experience)
    staff_display = entry.attending_staff_display
    return {
        'id': entry.id,
        'date': local_created.strftime('%Y-%m-%d'),
        'time': local_created.strftime('%H:%M'),
        'experience': exp_display,
        'rating': exp_display,
        'category': _CAT_DISPLAY.get(entry.category, entry.category),
        'category_value': entry.category,
        'sentiment': sentiment_display,
        'sentiment_value': sentiment_value,
        'status': _STATUS_DISPLAY.get(entry.status, entry.status),
        'status_value': entry.status,
        'comment': entry.comment,
        'name_of_client': entry.name_of_client or '',
        'client_type': entry.client_type or '',
        'staff': staff_display,
        'notes': notes,
        'status_history': status_history,
    }


def _experience_counts(qs):
    res = qs.aggregate(
        total=Count('id'),
        sa=Count('id', filter=Q(experience__in=[FeedbackEntry.STRONGLY_AGREE, 'vsat'])),
        a=Count('id', filter=Q(experience__in=[FeedbackEntry.AGREE, 'sat'])),
        nad=Count('id', filter=Q(experience=FeedbackEntry.NEITHER)),
        d=Count('id', filter=Q(experience=FeedbackEntry.DISAGREE)),
        sd=Count('id', filter=Q(experience__in=[FeedbackEntry.STRONGLY_DISAGREE, 'unsat'])),
        na=Count('id', filter=Q(experience=FeedbackEntry.NOT_APPLICABLE)),
    )
    total = res['total'] or 0
    sa = res['sa'] or 0
    a = res['a'] or 0
    nad = res['nad'] or 0
    d = res['d'] or 0
    sd = res['sd'] or 0
    na_count = res['na'] or 0
    return {
        'total': total,
        'strongly_agree': sa,
        'agree': a,
        'neither': nad,
        'disagree': d,
        'strongly_disagree': sd,
        'na': na_count,
    }


def _multi_period_experience_counts(qs, today_range, week_start, month_start):
    """Consolidates experience counts across all 4 timeframes into a single DB query."""
    def _cond(period_q, exp_q):
        return Count('id', filter=(period_q & exp_q) if period_q is not None else exp_q)

    def _period_aggs(prefix, period_q):
        sa_q = Q(experience__in=[FeedbackEntry.STRONGLY_AGREE, 'vsat'])
        a_q = Q(experience__in=[FeedbackEntry.AGREE, 'sat'])
        nad_q = Q(experience=FeedbackEntry.NEITHER)
        d_q = Q(experience=FeedbackEntry.DISAGREE)
        sd_q = Q(experience__in=[FeedbackEntry.STRONGLY_DISAGREE, 'unsat'])
        na_q = Q(experience=FeedbackEntry.NOT_APPLICABLE)
        tot_q = Count('id', filter=period_q) if period_q is not None else Count('id')
        return {
            f'{prefix}_total': tot_q,
            f'{prefix}_sa': _cond(period_q, sa_q),
            f'{prefix}_a': _cond(period_q, a_q),
            f'{prefix}_nad': _cond(period_q, nad_q),
            f'{prefix}_d': _cond(period_q, d_q),
            f'{prefix}_sd': _cond(period_q, sd_q),
            f'{prefix}_na': _cond(period_q, na_q),
        }

    aggs = {}
    aggs.update(_period_aggs('all', None))
    aggs.update(_period_aggs('today', Q(created_at__range=today_range)))
    aggs.update(_period_aggs('week', Q(created_at__gte=week_start)))
    aggs.update(_period_aggs('month', Q(created_at__gte=month_start)))

    res = qs.aggregate(**aggs)

    def _extract(prefix):
        return {
            'total': res[f'{prefix}_total'] or 0,
            'strongly_agree': res[f'{prefix}_sa'] or 0,
            'agree': res[f'{prefix}_a'] or 0,
            'neither': res[f'{prefix}_nad'] or 0,
            'disagree': res[f'{prefix}_d'] or 0,
            'strongly_disagree': res[f'{prefix}_sd'] or 0,
            'na': res[f'{prefix}_na'] or 0,
        }

    return {
        'all': _extract('all'),
        'today': _extract('today'),
        'week': _extract('week'),
        'month': _extract('month'),
    }


def _multi_period_sentiment_counts(qs, today_range, week_start, month_start):
    """Consolidates sentiment counts across all 4 timeframes into a single DB query."""
    def _cond(period_q, sent_val):
        sent_q = Q(sentiment=sent_val)
        return Count('id', filter=(period_q & sent_q) if period_q is not None else sent_q)

    def _period_aggs(prefix, period_q):
        return {
            f'{prefix}_pos': _cond(period_q, FeedbackEntry.POSITIVE),
            f'{prefix}_neu': _cond(period_q, FeedbackEntry.NEUTRAL),
            f'{prefix}_neg': _cond(period_q, FeedbackEntry.NEGATIVE),
        }

    aggs = {}
    aggs.update(_period_aggs('all', None))
    aggs.update(_period_aggs('today', Q(created_at__range=today_range)))
    aggs.update(_period_aggs('week', Q(created_at__gte=week_start)))
    aggs.update(_period_aggs('month', Q(created_at__gte=month_start)))

    res = qs.aggregate(**aggs)

    def _format(prefix):
        pos = res[f'{prefix}_pos'] or 0
        neu = res[f'{prefix}_neu'] or 0
        neg = res[f'{prefix}_neg'] or 0
        tot = pos + neu + neg
        def _pct(n):
            return round((n / tot) * 100) if tot else 0
        return {
            'total': tot,
            'positive': pos,
            'neutral': neu,
            'negative': neg,
            'positive_pct': _pct(pos),
            'neutral_pct': _pct(neu),
            'negative_pct': _pct(neg),
        }

    return {
        'all': _format('all'),
        'today': _format('today'),
        'week': _format('week'),
        'month': _format('month'),
    }


def _multi_period_report_data(qs, periods_map):
    """Consolidates metrics for daily, weekly, monthly, quarterly, and annual periods in 1 query."""
    aggs = {}
    for prefix, period_q in periods_map.items():
        aggs[f'{prefix}_total'] = Count('id', filter=period_q)
        aggs[f'{prefix}_sa'] = Count('id', filter=period_q & Q(experience__in=[FeedbackEntry.STRONGLY_AGREE, 'vsat']))
        aggs[f'{prefix}_a'] = Count('id', filter=period_q & Q(experience__in=[FeedbackEntry.AGREE, 'sat']))
        aggs[f'{prefix}_nad'] = Count('id', filter=period_q & Q(experience=FeedbackEntry.NEITHER))
        aggs[f'{prefix}_d'] = Count('id', filter=period_q & Q(experience=FeedbackEntry.DISAGREE))
        aggs[f'{prefix}_sd'] = Count('id', filter=period_q & Q(experience__in=[FeedbackEntry.STRONGLY_DISAGREE, 'unsat']))
        aggs[f'{prefix}_na'] = Count('id', filter=period_q & Q(experience=FeedbackEntry.NOT_APPLICABLE))
        aggs[f'{prefix}_compliment'] = Count('id', filter=period_q & Q(category='compliment'))
        aggs[f'{prefix}_suggestion'] = Count('id', filter=period_q & Q(category='suggestion'))
        aggs[f'{prefix}_complaint'] = Count('id', filter=period_q & Q(category='complaint'))
        aggs[f'{prefix}_concern'] = Count('id', filter=period_q & Q(category='concern'))

    res = qs.aggregate(**aggs)

    result = {}
    for prefix in periods_map:
        total = res[f'{prefix}_total'] or 0
        sa = res[f'{prefix}_sa'] or 0
        a = res[f'{prefix}_a'] or 0
        nad = res[f'{prefix}_nad'] or 0
        d = res[f'{prefix}_d'] or 0
        sd = res[f'{prefix}_sd'] or 0
        na_count = res[f'{prefix}_na'] or 0
        satisfaction = round((sa + a) / total * 100) if total else 0

        cat_counts = {
            'compliment': res[f'{prefix}_compliment'] or 0,
            'suggestion': res[f'{prefix}_suggestion'] or 0,
            'complaint': res[f'{prefix}_complaint'] or 0,
            'concern': res[f'{prefix}_concern'] or 0,
        }
        categorized = sum(cat_counts.values())

        result[prefix] = {
            'total': total,
            'strongly_agree': sa,
            'agree': a,
            'neither': nad,
            'disagree': d,
            'strongly_disagree': sd,
            'na': na_count,
            'vsat': sa,
            'sat': a,
            'neg': sd + d,
            'satisfaction': satisfaction,
            'categorized': categorized,
            'categories': cat_counts,
        }
    return result


@staff_required
@require_POST
def response_status_update(request, entry_id):
    entry = get_object_or_404(FeedbackEntry, pk=entry_id)
    try:
        payload = json.loads(request.body.decode('utf-8'))
    except json.JSONDecodeError:
        return JsonResponse({'ok': False, 'error': 'Invalid request.'}, status=400)

    status = payload.get('status')
    valid_statuses = {choice[0] for choice in FeedbackEntry.STATUS_CHOICES}
    if status not in valid_statuses:
        return JsonResponse({'ok': False, 'error': 'Invalid status.'}, status=400)

    status_display = dict(FeedbackEntry.STATUS_CHOICES)
    old_status = entry.status
    history_entry = None

    if status != old_status:
        entry.status = status
        entry.save(update_fields=['status', 'updated_at'])
        log_feedback_status_change(request.user, entry, old_status, status)
        history_entry = {
            'old': status_display.get(old_status, old_status),
            'new': status_display.get(status, status),
            'by': request.user.get_full_name() or request.user.username,
            'at': timezone.localtime(timezone.now()).strftime('%b %d, %Y %I:%M %p'),
        }

    return JsonResponse({
        'ok': True,
        'status': entry.get_status_display(),
        'status_value': entry.status,
        'updated_at': timezone.localtime(entry.updated_at).strftime('%b %d, %Y %I:%M %p'),
        'history_entry': history_entry,
    })


@staff_required
@require_POST
def response_category_update(request, entry_id):
    entry = get_object_or_404(FeedbackEntry, pk=entry_id)
    try:
        payload = json.loads(request.body.decode('utf-8'))
    except json.JSONDecodeError:
        return JsonResponse({'ok': False, 'error': 'Invalid request.'}, status=400)

    category = payload.get('category')
    valid_categories = {choice[0] for choice in FeedbackEntry.CATEGORY_CHOICES}
    if category and category not in valid_categories:
        return JsonResponse({'ok': False, 'error': 'Invalid category.'}, status=400)

    old_category = entry.category
    entry.category = category or ''
    if entry.category != old_category:
        entry.save(update_fields=['category', 'updated_at'])
        log_admin_event(
            request.user,
            entry,
            CHANGE,
            f'category:{old_category or ""}|{entry.category or ""}',
        )

    return JsonResponse({
        'ok': True,
        'category': entry.get_category_display(),
        'category_value': entry.category,
    })


@staff_required
def responses_count(request):
    return JsonResponse({
        'ok': True,
        'count': FeedbackEntry.objects.count(),
    })


@staff_required
@require_POST
def responses_delete(request):
    try:
        payload = json.loads(request.body.decode('utf-8'))
    except (json.JSONDecodeError, UnicodeDecodeError, AttributeError):
        return JsonResponse({'ok': False, 'error': 'Invalid request payload.'}, status=400)

    raw_ids = payload.get('ids', [])
    if not isinstance(raw_ids, list) or not raw_ids:
        return JsonResponse({'ok': False, 'error': 'No response IDs provided.'}, status=400)

    clean_ids = []
    for item in raw_ids:
        try:
            clean_ids.append(int(item))
        except (ValueError, TypeError):
            continue

    if not clean_ids:
        return JsonResponse({'ok': False, 'error': 'No valid response IDs provided.'}, status=400)

    entries = list(FeedbackEntry.objects.filter(id__in=clean_ids))
    if not entries:
        return JsonResponse({'ok': False, 'error': 'Selected responses not found or already deleted.'}, status=404)

    deleted_count = len(entries)
    deleted_ids = [entry.id for entry in entries]

    with transaction.atomic():
        for entry in entries:
            log_admin_event(
                request.user,
                entry,
                DELETION,
                f'Deleted feedback response #{entry.id}',
            )
        FeedbackEntry.objects.filter(id__in=deleted_ids).delete()

    counts = _experience_counts(FeedbackEntry.objects.all())
    noun = 'response' if deleted_count == 1 else 'responses'

    return JsonResponse({
        'ok': True,
        'message': f'Successfully deleted {deleted_count} {noun}.',
        'deleted_ids': deleted_ids,
        'deleted_count': deleted_count,
        'counts': counts,
    })


@superuser_required
def sentiment_analysis(request):
    entries = FeedbackEntry.objects.all()
    now = timezone.localtime(timezone.now())
    today = now.date()
    week_start = now - timedelta(days=7)
    month_start = now - timedelta(days=30)
    today_start = timezone.make_aware(datetime.combine(today, time.min))
    today_end = timezone.make_aware(datetime.combine(today, time.max))

    filter_data = _multi_period_sentiment_counts(entries, (today_start, today_end), week_start, month_start)
    all_counts = filter_data['all']

    trend_counts = (
        entries.exclude(sentiment=FeedbackEntry.PENDING)
        .annotate(local_date=TruncDate('created_at'))
        .values('local_date', 'sentiment')
        .annotate(count=Count('id'))
    )
    trend_data_map = defaultdict(lambda: defaultdict(int))
    for row in trend_counts:
        if row['local_date']:
            trend_data_map[row['local_date']][row['sentiment']] = row['count']

    if trend_data_map:
        earliest_date = min(trend_data_map.keys())
        start_date = min(earliest_date, today - timedelta(days=6))
        total_days = (today - start_date).days + 1
        trend_dates = [start_date + timedelta(days=i) for i in range(total_days)]
    else:
        trend_dates = [today - timedelta(days=i) for i in range(6, -1, -1)]

    trend_labels = [f'{day:%b} {day.day}' for day in trend_dates]

    def trend_for(sentiment):
        return [trend_data_map[day][sentiment] for day in trend_dates]

    context = {
        'total': all_counts['total'],
        'positive': all_counts['positive'],
        'neutral': all_counts['neutral'],
        'negative': all_counts['negative'],
        'positive_pct': all_counts['positive_pct'],
        'neutral_pct': all_counts['neutral_pct'],
        'negative_pct': all_counts['negative_pct'],
        'filter_data': filter_data,

        'trend_labels': trend_labels,
        'trend_dates': [day.isoformat() for day in trend_dates],
        'trend_positive': trend_for(FeedbackEntry.POSITIVE),
        'trend_neutral': trend_for(FeedbackEntry.NEUTRAL),
        'trend_negative': trend_for(FeedbackEntry.NEGATIVE),
    }
    return render(request, 'feedback_admin/sentiment_analysis.html', context)


@staff_required
def reports(request):
    now = timezone.localtime(timezone.now())
    today = now.date()
    week_start = now - timedelta(days=7)
    month_start = now - timedelta(days=30)
    quarter_start = now - timedelta(days=90)
    year_start = now - timedelta(days=365)

    today_start = timezone.make_aware(datetime.combine(today, time.min))
    today_end = timezone.make_aware(datetime.combine(today, time.max))

    periods_map = {
        'daily': Q(created_at__range=(today_start, today_end)),
        'weekly': Q(created_at__gte=week_start),
        'monthly': Q(created_at__gte=month_start),
        'quarterly': Q(created_at__gte=quarter_start),
        'annual': Q(created_at__gte=year_start),
    }
    period_stats = _multi_period_report_data(FeedbackEntry.objects.all(), periods_map)
    daily_data = period_stats['daily']
    weekly_data = period_stats['weekly']
    monthly_data = period_stats['monthly']
    quarterly_data = period_stats['quarterly']
    annual_data = period_stats['annual']

    daily_qs = FeedbackEntry.objects.filter(created_at__range=(today_start, today_end))
    weekly_qs = FeedbackEntry.objects.filter(created_at__gte=week_start)
    monthly_qs = FeedbackEntry.objects.filter(created_at__gte=month_start)
    quarterly_qs = FeedbackEntry.objects.filter(created_at__gte=quarter_start)
    annual_qs = FeedbackEntry.objects.filter(created_at__gte=year_start)

    def _hourly_trend(qs):
        """Buckets entries by hour-of-day (0–23) — for a single day's queryset."""
        counts_qs = qs.annotate(hour=ExtractHour('created_at')).values('hour').annotate(count=Count('id'))
        hour_map = {row['hour']: row['count'] for row in counts_qs if row['hour'] is not None}
        counts = [hour_map.get(h, 0) for h in range(24)]
        labels = [datetime(2000, 1, 1, hour).strftime('%I %p').lstrip('0') for hour in range(24)]
        return labels, counts

    def _daily_trend(qs, start_date, end_date):
        """Buckets entries by calendar day across an inclusive date range."""
        counts_qs = qs.annotate(local_date=TruncDate('created_at')).values('local_date').annotate(count=Count('id'))
        counts = {row['local_date']: row['count'] for row in counts_qs if row['local_date']}

        days, cursor = [], start_date
        while cursor <= end_date:
            days.append(cursor)
            cursor += timedelta(days=1)

        labels = [f'{day:%b} {day.day}' for day in days]
        data = [counts.get(day, 0) for day in days]
        return labels, data

    def _weekly_trend(qs, start_date, end_date):
        """Buckets entries by ISO week (Mon–Sun) across an inclusive date range."""
        counts_qs = qs.annotate(local_week=TruncWeek('created_at')).values('local_week').annotate(count=Count('id'))
        counts = {row['local_week'].date() if hasattr(row['local_week'], 'date') else row['local_week']: row['count'] for row in counts_qs if row['local_week']}

        first_week = start_date - timedelta(days=start_date.weekday())
        last_week = end_date - timedelta(days=end_date.weekday())

        weeks, cursor = [], first_week
        while cursor <= last_week:
            weeks.append(cursor)
            cursor += timedelta(days=7)

        labels = [f'{week:%b} {week.day}' for week in weeks]
        data = [counts.get(week, 0) for week in weeks]
        return labels, data

    def _monthly_trend(qs, months_back):
        """Buckets entries by calendar month for the trailing `months_back` months."""
        counts_qs = qs.annotate(local_month=TruncMonth('created_at')).values('local_month').annotate(count=Count('id'))
        counts = {row['local_month'].date().replace(day=1) if hasattr(row['local_month'], 'date') else row['local_month']: row['count'] for row in counts_qs if row['local_month']}

        months, cursor = [], timezone.localdate().replace(day=1)
        for _ in range(months_back):
            months.append(cursor)
            cursor = (cursor - timedelta(days=1)).replace(day=1)
        months.reverse()

        labels = [f'{month:%b %Y}' for month in months]
        data = [counts.get(month, 0) for month in months]
        return labels, data

    # ── Response trend, bucketed at a granularity that fits each period ──
    daily_data['trendLabels'], daily_data['trendData'] = _hourly_trend(daily_qs)
    weekly_data['trendLabels'], weekly_data['trendData'] = _daily_trend(
        weekly_qs, (now - timedelta(days=6)).date(), today)
    monthly_data['trendLabels'], monthly_data['trendData'] = _daily_trend(
        monthly_qs, (now - timedelta(days=29)).date(), today)
    quarterly_data['trendLabels'], quarterly_data['trendData'] = _weekly_trend(
        quarterly_qs, (now - timedelta(days=89)).date(), today)
    annual_data['trendLabels'], annual_data['trendData'] = _monthly_trend(annual_qs, 12)

    # For now, let's just use the basic stats in context
    context = {
        'daily_total': daily_data['total'],
        'daily_vsat': daily_data['vsat'],
        'daily_sat': daily_data['sat'],
        'daily_neg': daily_data['neg'],
        'daily_satisfaction': daily_data['satisfaction'],

        'weekly_total': weekly_data['total'],
        'weekly_vsat': weekly_data['vsat'],
        'weekly_sat': weekly_data['sat'],
        'weekly_neg': weekly_data['neg'],
        'weekly_satisfaction': weekly_data['satisfaction'],

        'monthly_total': monthly_data['total'],
        'monthly_vsat': monthly_data['vsat'],
        'monthly_sat': monthly_data['sat'],
        'monthly_neg': monthly_data['neg'],
        'monthly_satisfaction': monthly_data['satisfaction'],

        'quarterly_total': quarterly_data['total'],
        'quarterly_vsat': quarterly_data['vsat'],
        'quarterly_sat': quarterly_data['sat'],
        'quarterly_neg': quarterly_data['neg'],
        'quarterly_satisfaction': quarterly_data['satisfaction'],

        'annual_total': annual_data['total'],
        'annual_vsat': annual_data['vsat'],
        'annual_sat': annual_data['sat'],
        'annual_neg': annual_data['neg'],
        'annual_satisfaction': annual_data['satisfaction'],

        # Pass the whole structured object for JS
        'report_json': {
            'daily': daily_data,
            'weekly': weekly_data,
            'monthly': monthly_data,
            'quarterly': quarterly_data,
            'annual': annual_data,
        }
    }

    return render(request, 'feedback_admin/reports.html', context)


@staff_required
def export_report_excel(request):
    """Generate and return an .xlsx report for the selected period."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    period = request.GET.get('period', 'daily')
    now = timezone.localtime(timezone.now())
    today = now.date()

    # Period date ranges — mirrors the reports() view logic
    today_start = timezone.make_aware(datetime.combine(today, time.min))
    today_end = timezone.make_aware(datetime.combine(today, time.max))

    period_ranges = {
        'daily': (today_start, today_end),
        'weekly': (now - timedelta(days=7), now),
        'monthly': (now - timedelta(days=30), now),
        'quarterly': (now - timedelta(days=90), now),
        'annual': (now - timedelta(days=365), now),
    }
    if period not in period_ranges:
        period = 'daily'

    start, end = period_ranges[period]
    qs = FeedbackEntry.objects.filter(created_at__range=(start, end))

    # Compute summary stats (reuse the existing helper)
    periods_map = {period: Q(created_at__range=(start, end))}
    stats = _multi_period_report_data(FeedbackEntry.objects.all(), periods_map)[period]

    # ── Styles ────────────────────────────────────────────────────────────
    header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='0F5A2B', end_color='0F5A2B', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )
    label_font = Font(name='Calibri', bold=True, size=11)
    value_font = Font(name='Calibri', size=11)
    title_font = Font(name='Calibri', bold=True, size=14, color='0F5A2B')

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Summary'
    ws.sheet_properties.tabColor = '169651'

    # Title
    ws.merge_cells('A1:B1')
    ws['A1'] = f'PhilHealth CSM Report — {period.title()}'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(vertical='center')
    ws['A2'] = f'Generated: {now.strftime("%B %d, %Y at %I:%M %p")}'
    ws['A2'].font = Font(name='Calibri', size=10, italic=True, color='475569')
    ws.append([])  # blank row

    # Overview metrics
    overview_rows = [
        ('Total Responses', stats['total']),
        ('Satisfaction Rate', f"{stats['satisfaction']}%"),
        ('', ''),
        ('SQD Distribution', ''),
        ('Strongly Agree', stats['strongly_agree']),
        ('Agree', stats['agree']),
        ('Neither Agree nor Disagree', stats['neither']),
        ('Disagree', stats['disagree']),
        ('Strongly Disagree', stats['strongly_disagree']),
        ('Not Applicable', stats['na']),
        ('', ''),
        ('Feedback Categories', ''),
        ('Compliments', stats['categories']['compliment']),
        ('Suggestions', stats['categories']['suggestion']),
        ('Complaints', stats['categories']['complaint']),
        ('Service Concerns', stats['categories']['concern']),
        ('Total Categorized', stats['categorized']),
    ]
    for label, value in overview_rows:
        ws.append([label, value])
        row_num = ws.max_row
        ws.cell(row=row_num, column=1).font = label_font if label and value == '' else value_font
        ws.cell(row=row_num, column=2).font = value_font
        if label and value != '':
            ws.cell(row=row_num, column=1).font = label_font
        for col in (1, 2):
            ws.cell(row=row_num, column=col).border = thin_border

    # Section headers (SQD Distribution, Feedback Categories) get bold styling
    for row_idx in range(1, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=1)
        if cell.value in ('SQD Distribution', 'Feedback Categories'):
            cell.font = Font(name='Calibri', bold=True, size=11, color='0F5A2B')

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 18

    # ── Sheet 2: Responses ────────────────────────────────────────────────
    ws2 = wb.create_sheet('Responses')
    ws2.sheet_properties.tabColor = '23A455'

    response_headers = [
        'ID', 'Date', 'Time', 'Client Name', 'Age', 'Sex', 'Client Type',
        'CC1', 'CC2', 'CC3',
        'SQD0', 'SQD1', 'SQD2', 'SQD3', 'SQD4', 'SQD5', 'SQD6', 'SQD7', 'SQD8',
        'Experience', 'Category', 'Sentiment', 'Status',
        'Staff Assisted', 'Comment', 'Suggestions', 'Commendation',
    ]

    # Write header row
    for col_idx, header_text in enumerate(response_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Write data rows
    EXPERIENCE_MAP = dict(FeedbackEntry.EXPERIENCE_CHOICES)
    SENTIMENT_MAP = dict(FeedbackEntry.SENTIMENT_CHOICES)
    CATEGORY_MAP = dict(FeedbackEntry.CATEGORY_CHOICES)
    STATUS_MAP = dict(FeedbackEntry.STATUS_CHOICES)

    entries = qs.select_related('staff_assisted').order_by('-created_at')
    for row_idx, entry in enumerate(entries, 2):
        local_dt = timezone.localtime(entry.created_at) if entry.created_at else None
        row_data = [
            entry.pk,
            local_dt.strftime('%Y-%m-%d') if local_dt else '',
            local_dt.strftime('%I:%M %p') if local_dt else '',
            entry.name_of_client,
            entry.age,
            entry.sex,
            entry.client_type,
            entry.cc1, entry.cc2, entry.cc3,
            entry.sqd0, entry.sqd1, entry.sqd2, entry.sqd3,
            entry.sqd4, entry.sqd5, entry.sqd6, entry.sqd7, entry.sqd8,
            EXPERIENCE_MAP.get(entry.experience, entry.experience),
            CATEGORY_MAP.get(entry.category, entry.category),
            SENTIMENT_MAP.get(entry.sentiment, entry.sentiment),
            STATUS_MAP.get(entry.status, entry.status),
            entry.attending_staff_display,
            entry.comment,
            entry.comments_suggestions,
            entry.commendation,
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.font = value_font
            cell.border = thin_border

    # Auto-size key columns (approximate widths)
    col_widths = {
        'A': 8, 'B': 12, 'C': 10, 'D': 22, 'E': 6, 'F': 8, 'G': 14,
        'H': 6, 'I': 6, 'J': 6,
        'K': 6, 'L': 6, 'M': 6, 'N': 6, 'O': 6, 'P': 6, 'Q': 6, 'R': 6, 'S': 6,
        'T': 18, 'U': 14, 'V': 12, 'W': 12,
        'X': 22, 'Y': 36, 'Z': 36, 'AA': 36,
    }
    for col_letter, width in col_widths.items():
        ws2.column_dimensions[col_letter].width = width

    # Freeze the header row
    ws2.freeze_panes = 'A2'

    # ── Return the file ───────────────────────────────────────────────────
    period_labels = {
        'daily': 'Daily', 'weekly': 'Weekly', 'monthly': 'Monthly',
        'quarterly': 'Quarterly', 'annual': 'Annual',
    }
    filename = f'PhilHealth-Report-{period_labels[period]}-{today.isoformat()}.xlsx'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@superuser_required
def activity_log(request):
    """
    Government-facing audit trail sourced entirely from Django's built-in
    django_admin_log table.
    """
    logs = list(
        LogEntry.objects
        .select_related('user', 'content_type')
        .order_by('-action_time')
    )

    feedback_ct = _feedback_content_type()
    feedback_ct_id = feedback_ct.pk
    valid_pks = [
        int(log.object_id)
        for log in logs
        if log.content_type_id == feedback_ct_id and log.object_id and log.object_id.isdigit()
    ]
    feedback_entries = {e.pk: e for e in FeedbackEntry.objects.filter(pk__in=valid_pks)}

    rows = []
    for log in logs:
        is_fb = log.content_type_id == feedback_ct_id
        entry = feedback_entries.get(int(log.object_id)) if is_fb and log.object_id and log.object_id.isdigit() else None
        rows.append(_format_audit_row(log, entry))

    context = {
        'logs_data': rows,
        'total_actions': len(rows),
    }
    return render(request, 'feedback_admin/activity_log.html', context)


# ── user management ───────────────────────────────────────────────────

@superuser_required
def users(request):
    all_users = User.objects.all().order_by('-date_joined').prefetch_related('groups', 'user_permissions')
    all_groups = Group.objects.annotate(member_count=Count('user')).prefetch_related('permissions')
    all_perms = Permission.objects.select_related('content_type').order_by('content_type__app_label', 'codename')

    user_counts = User.objects.aggregate(
        total=Count('id'),
        active=Count('id', filter=Q(is_active=True)),
        staff=Count('id', filter=Q(is_staff=True)),
    )
    evaluated_groups = list(all_groups)

    context = {
        'users': all_users,
        'groups': evaluated_groups,
        'permissions': all_perms,
        'total_users': user_counts['total'] or 0,
        'active_users': user_counts['active'] or 0,
        'staff_users': user_counts['staff'] or 0,
        'total_groups': len(evaluated_groups),
    }
    return render(request, 'feedback_admin/users.html', context)


@superuser_required
@require_POST
def user_add(request):
    def err(msg):
        if _is_ajax(request):
            return JsonResponse({'ok': False, 'error': msg})
        messages.error(request, msg)
        return redirect('users')

    username = request.POST.get('username', '').strip()
    if not username:
        return err('Username is required.')
    if User.objects.filter(username=username).exists():
        return err(f'Username "{username}" is already taken.')

    pw1 = request.POST.get('password1', '')
    pw2 = request.POST.get('password2', '')
    if pw1 != pw2:
        return err('Passwords do not match.')

    tmp_user = User(
        username=username,
        email=request.POST.get('email', '').strip(),
        first_name=request.POST.get('first_name', '').strip(),
        last_name=request.POST.get('last_name', '').strip(),
    )
    try:
        validate_password(pw1, user=tmp_user)
    except ValidationError as e:
        return err(' '.join(e.messages))

    has_usable_password = request.POST.get('has_usable_password', 'on') == 'on'

    user = User.objects.create(
        username=username,
        email=tmp_user.email,
        first_name=tmp_user.first_name,
        last_name=tmp_user.last_name,
        is_active='is_active' in request.POST,
        is_staff='is_staff' in request.POST,
        is_superuser='is_superuser' in request.POST,
    )
    if has_usable_password:
        user.set_password(pw1)
    else:
        user.set_unusable_password()
    user.save()

    group_ids = request.POST.getlist('groups')
    if group_ids:
        user.groups.set(Group.objects.filter(id__in=group_ids))
    perm_ids = request.POST.getlist('user_permissions')
    if perm_ids:
        user.user_permissions.set(Permission.objects.filter(id__in=perm_ids))

    log_admin_event(request.user, user, ADDITION, f'Created user "{username}"')

    msg = f'User "{username}" created successfully.'
    if _is_ajax(request):
        return JsonResponse({'ok': True, 'message': msg})
    messages.success(request, msg)
    return redirect('users')


@staff_required
@require_POST
def user_edit(request, user_id):
    user = get_object_or_404(User, pk=user_id)
    if not (request.user.is_superuser or request.user == user):
        if _is_ajax(request):
            return JsonResponse({'ok': False, 'error': 'Access restricted to administrators.'}, status=403)
        messages.error(request, 'Access restricted to administrators.')
        return redirect('dashboard')

    original = {
        'username': user.username,
        'email': user.email,
        'first_name': user.first_name,
        'last_name': user.last_name,
        'is_active': user.is_active,
        'is_staff': user.is_staff,
        'is_superuser': user.is_superuser,
        'groups': list(user.groups.values_list('id', flat=True)),
        'perms': list(user.user_permissions.values_list('id', flat=True)),
        'has_password': user.has_usable_password(),
    }

    def err(msg):
        if _is_ajax(request):
            return JsonResponse({'ok': False, 'error': msg})
        messages.error(request, msg)
        return redirect('users')

    username = request.POST.get('username', '').strip()
    if not username:
        return err('Username is required.')
    if User.objects.filter(username=username).exclude(pk=user_id).exists():
        return err(f'Username "{username}" is already taken.')

    user.username = username
    user.email = request.POST.get('email', '').strip()
    user.first_name = request.POST.get('first_name', '').strip()
    user.last_name = request.POST.get('last_name', '').strip()

    if request.user.is_superuser:
        if user == request.user:
            user.is_active = True
        else:
            user.is_active = 'is_active' in request.POST

        user.is_staff = 'is_staff' in request.POST
        user.is_superuser = 'is_superuser' in request.POST

    pw1 = request.POST.get('password1', '')
    pw2 = request.POST.get('password2', '')
    if pw1:
        if user == request.user:
            current_pw = request.POST.get('current_password', '')
            if not current_pw:
                return err('Current password is required to change your password.')
            if not user.check_password(current_pw):
                return err('Current password is incorrect.')

        if pw1 != pw2:
            return err('Passwords do not match.')
        try:
            validate_password(pw1, user=user)
        except ValidationError as e:
            return err(' '.join(e.messages))
        user.set_password(pw1)
        if user == request.user:
            from django.contrib.auth import update_session_auth_hash
            update_session_auth_hash(request, user)

    has_usable_password = request.POST.get('has_usable_password', 'on') == 'on'
    if not pw1:
        if has_usable_password and not user.has_usable_password():
            pass
        elif not has_usable_password:
            user.set_unusable_password()

    user.save()
    if request.user.is_superuser:
        user.groups.set(Group.objects.filter(id__in=request.POST.getlist('groups')))
        user.user_permissions.set(Permission.objects.filter(id__in=request.POST.getlist('user_permissions')))

    changed_fields = []
    for key, label in [
        ('username', 'username'),
        ('email', 'email'),
        ('first_name', 'first name'),
        ('last_name', 'last name'),
    ]:
        if original[key] != getattr(user, key):
            changed_fields.append(label)
    if original['is_active'] != user.is_active:
        changed_fields.append('status')
    if original['is_staff'] != user.is_staff:
        changed_fields.append('staff role')
    if original['is_superuser'] != user.is_superuser:
        changed_fields.append('superuser role')
    if original['groups'] != list(user.groups.values_list('id', flat=True)):
        changed_fields.append('groups')
    if original['perms'] != list(user.user_permissions.values_list('id', flat=True)):
        changed_fields.append('permissions')
    if original['has_password'] != user.has_usable_password() or pw1:
        changed_fields.append('password')

    if changed_fields:
        target = 'admin profile' if user == request.user else f'user "{user.username}"'
        log_admin_event(
            request.user,
            user,
            CHANGE,
            f'Updated {target}: {", ".join(changed_fields)}',
        )

    msg = f'User "{user.username}" updated successfully.'
    if _is_ajax(request):
        return JsonResponse({'ok': True, 'message': msg})
    messages.success(request, msg)
    return redirect('users')


@superuser_required
@require_POST
def user_delete(request, user_id):
    user = get_object_or_404(User, pk=user_id)
    if user == request.user:
        if _is_ajax(request):
            return JsonResponse({'ok': False, 'error': 'You cannot delete your own account.'})
        messages.error(request, "You cannot delete your own account.")
        return redirect('users')
    username = user.username
    log_admin_event(request.user, user, DELETION, f'Deleted user "{username}"')
    user.delete()

    msg = f'User "{username}" deleted.'
    if _is_ajax(request):
        return JsonResponse({'ok': True, 'message': msg})
    messages.success(request, msg)
    return redirect('users')


@superuser_required
@require_POST
def user_toggle_active(request, user_id):
    user = get_object_or_404(User, pk=user_id)
    if user == request.user:
        messages.error(request, "You cannot deactivate your own account.")
        return redirect('users')
    user.is_active = not user.is_active
    user.save(update_fields=['is_active'])
    state = 'activated' if user.is_active else 'deactivated'
    log_admin_event(request.user, user, CHANGE, f'User "{user.username}" {state}')
    messages.success(request, f'User "{user.username}" {state}.')
    return redirect('users')


@superuser_required
@require_POST
def group_add(request):
    def err(msg):
        if _is_ajax(request):
            return JsonResponse({'ok': False, 'error': msg})
        messages.error(request, msg)
        return redirect('users')

    name = request.POST.get('name', '').strip()
    if not name:
        return err('Group name is required.')
    if Group.objects.filter(name=name).exists():
        return err(f'Group "{name}" already exists.')

    group = Group.objects.create(name=name)
    perm_ids = request.POST.getlist('permissions')
    if perm_ids:
        group.permissions.set(Permission.objects.filter(id__in=perm_ids))

    log_admin_event(request.user, group, ADDITION, f'Created group "{name}"')

    msg = f'Group "{name}" created.'
    if _is_ajax(request):
        return JsonResponse({'ok': True, 'message': msg})
    messages.success(request, msg)
    return redirect('users')


@superuser_required
@require_POST
def group_edit(request, group_id):
    group = get_object_or_404(Group, pk=group_id)
    original_name = group.name
    original_perms = list(group.permissions.values_list('id', flat=True))

    def err(msg):
        if _is_ajax(request):
            return JsonResponse({'ok': False, 'error': msg})
        messages.error(request, msg)
        return redirect('users')

    name = request.POST.get('name', '').strip()
    if not name:
        return err('Group name is required.')
    if Group.objects.filter(name=name).exclude(pk=group_id).exists():
        return err(f'Group "{name}" already exists.')

    group.name = name
    group.save()
    group.permissions.set(Permission.objects.filter(id__in=request.POST.getlist('permissions')))

    if original_name != group.name or original_perms != list(group.permissions.values_list('id', flat=True)):
        log_admin_event(request.user, group, CHANGE, f'Updated group "{group.name}"')

    msg = f'Group "{name}" updated.'
    if _is_ajax(request):
        return JsonResponse({'ok': True, 'message': msg})
    messages.success(request, msg)
    return redirect('users')


@superuser_required
@require_POST
def group_delete(request, group_id):
    group = get_object_or_404(Group, pk=group_id)
    name = group.name
    log_admin_event(request.user, group, DELETION, f'Deleted group "{name}"')
    group.delete()

    msg = f'Group "{name}" deleted.'
    if _is_ajax(request):
        return JsonResponse({'ok': True, 'message': msg})
    messages.success(request, msg)
    return redirect('users')


# ── LOGIN ─────────────────────────────────────────────────────────────
def admin_login(request):
    if request.user.is_authenticated:
        if request.user.is_staff or request.user.is_superuser:
            return redirect('dashboard')
        logout(request)

    form = AuthenticationForm(request, data=request.POST or None)

    if request.method == 'POST':
        if form.is_valid():
            user = form.get_user()
            if not (user.is_staff or user.is_superuser):
                messages.error(request, 'Access denied. Staff privileges are required.')
                return render(request, 'feedback_admin/login.html', {'form': form})
            login(request, user)
            log_admin_event(user, user, ADDITION, 'Logged in')
            return redirect('dashboard')
        else:
            username = request.POST.get('username', '').strip()
            password = request.POST.get('password', '')
            try:
                potential_user = User.objects.get(username=username)
                if potential_user.check_password(password):
                    if not potential_user.is_active:
                        messages.error(request, 'This account has been deactivated. Please contact your system administrator.')
                    elif not (potential_user.is_staff or potential_user.is_superuser):
                        messages.error(request, 'Access denied. Staff privileges are required.')
            except User.DoesNotExist:
                pass

    return render(request, 'feedback_admin/login.html', {'form': form})

# ── LOGOUT ────────────────────────────────────────────────────────────
@require_POST
def admin_logout(request):
    if request.user.is_authenticated:
        log_admin_event(request.user, request.user, CHANGE, 'Logged out')
    logout(request)
    messages.success(request, 'You have been signed out.')
    return redirect('admin_login')

@superuser_required
def settings_page(request):
    try:
        backups = list_backups()
        backup_error = None
    except Exception as e:
        backups = []
        backup_error = str(e)

    context = {
        'feedback_config': FeedbackConfiguration.get_solo(),
        'backups': backups,
        'backup_error': backup_error,
    }
    return render(request, 'feedback_admin/settings.html', context)


@superuser_required
@require_POST
def update_survey_settings(request):
    config = FeedbackConfiguration.get_solo()
    old_survey_enabled = config.survey_enabled
    survey_enabled = request.POST.get('survey_enabled') == 'on'
    offline_message = request.POST.get('survey_offline_message', '').strip()
    if not offline_message:
        offline_message = 'Ang feedback system ay pansamantalang hindi available. Pakisubukan muli mamaya.'

    config.survey_enabled = survey_enabled
    config.survey_offline_message = offline_message
    config.save(update_fields=['survey_enabled', 'survey_offline_message', 'updated_at'])

    if old_survey_enabled != survey_enabled:
        log_admin_event(
            request.user,
            config,
            CHANGE,
            f'Survey availability {"enabled" if survey_enabled else "disabled"}',
        )

    state = 'enabled' if survey_enabled else 'disabled'
    return JsonResponse({
        'ok': True,
        'message': f'Survey availability {state}. Citizens {"can now submit feedback via the public-facing form" if survey_enabled else "will see the offline notice when visiting the public portal"}.',
        'survey_enabled': survey_enabled,
        'survey_offline_message': offline_message,
    })


from feedback.services import reanalyze_pending_entries


@superuser_required
@require_POST
def update_sentiment_settings(request):
    old_value = FeedbackConfiguration.get_solo().auto_analysis_enabled
    auto_analysis_enabled = request.POST.get('auto_analysis_enabled') == 'on'
    config = FeedbackConfiguration.get_solo()
    config.auto_analysis_enabled = auto_analysis_enabled
    config.save(update_fields=['auto_analysis_enabled', 'updated_at'])

    if old_value != auto_analysis_enabled:
        log_admin_event(
            request.user,
            config,
            CHANGE,
            f'Auto-analysis {"enabled" if auto_analysis_enabled else "disabled"}',
        )

    state = 'enabled' if auto_analysis_enabled else 'disabled'
    return JsonResponse({
        'ok': True,
        'message': f'Auto-analysis {state}. New feedback will {"be analyzed for sentiment immediately" if auto_analysis_enabled else "stay pending until you re-enable auto-analysis or run re-analysis"}.',
        'auto_analysis_enabled': auto_analysis_enabled,
    })


@superuser_required
@require_POST
def reanalyze_sentiment_view(request):
    total, processed = reanalyze_pending_entries(force=False)
    config = FeedbackConfiguration.get_solo()
    log_admin_event(
        request.user,
        config,
        CHANGE,
        f'Batch re-analysis: scanned {total}, updated {processed}',
    )
    if total == 0:
        message = 'Nothing to process — no pending entries with comments found.'
    else:
        message = f'Re-analysis complete: {processed} of {total} pending entries categorized.'
    return JsonResponse({'ok': True, 'message': message, 'processed': processed, 'total': total})


@superuser_required
@require_POST
def update_notification_settings(request):
    config = FeedbackConfiguration.get_solo()

    daily_summary_enabled = request.POST.get('daily_summary_enabled') == 'on'
    notification_email = request.POST.get('notification_email', '').strip()
    daily_summary_time = request.POST.get('daily_summary_time', '17:30').strip()

    if notification_email:
        try:
            validate_email(notification_email)
        except ValidationError:
            return JsonResponse({
                'ok': False,
                'error': 'Please enter a valid notification email address.'
            }, status=400)

    config.daily_summary_enabled = daily_summary_enabled
    config.notification_email = notification_email
    config.daily_summary_time = daily_summary_time or '17:30'
    config.save(update_fields=[
        'daily_summary_enabled',
        'notification_email',
        'daily_summary_time',
        'updated_at',
    ])

    log_admin_event(
        request.user,
        config,
        CHANGE,
        f'Notification settings updated: daily summary {"enabled" if daily_summary_enabled else "disabled"}, email="{notification_email}"',
    )

    return JsonResponse({
        'ok': True,
        'message': 'Notification settings saved successfully.',
        'daily_summary_enabled': daily_summary_enabled,
        'notification_email': notification_email,
        'daily_summary_time': config.daily_summary_time,
    })


@superuser_required
@require_POST
def send_summary_now(request):
    config = FeedbackConfiguration.get_solo()
    target_email = request.POST.get('email', '').strip() or config.notification_email or request.user.email
    is_test = request.POST.get('is_test') == 'true'

    if not target_email:
        return JsonResponse({
            'ok': False,
            'error': 'No notification email address found. Please specify an email address first.'
        }, status=400)

    try:
        validate_email(target_email)
    except ValidationError:
        return JsonResponse({
            'ok': False,
            'error': f'"{target_email}" is not a valid email address.'
        }, status=400)

    base_url = request.build_absolute_uri('/')
    result = send_daily_summary_email(
        recipient_email=target_email,
        force=True,
        is_test=is_test,
        base_url=base_url
    )

    if result.get('ok'):
        action_desc = 'Test daily summary' if is_test else 'Daily summary'
        log_admin_event(
            request.user,
            config,
            CHANGE,
            f'{action_desc} email dispatched manually to {target_email}',
        )
        return JsonResponse({
            'ok': True,
            'message': f'Daily summary successfully dispatched to {target_email}.',
            'recipient': target_email,
        })
    else:
        return JsonResponse({
            'ok': False,
            'error': result.get('message', 'Failed to dispatch email.')
        }, status=500)


send_test_daily_summary = send_summary_now


@csrf_exempt
def cron_daily_summary(request):
    """
    Automated daily feedback summary endpoint for Vercel Cron Jobs,
    scheduled workflows, or external webhook calls.
    Supports GET (standard Vercel Cron) and POST.
    Protected by CRON_SECRET if configured in environment variables.
    """
    configured_secret = (getattr(settings, 'CRON_SECRET', '') or os.environ.get('CRON_SECRET', '')).strip()
    if configured_secret:
        auth_header = request.headers.get('Authorization', '')
        secret_param = request.GET.get('secret') or request.GET.get('key')
        expected_bearer = f"Bearer {configured_secret}"
        if auth_header != expected_bearer and secret_param != configured_secret:
            return JsonResponse({
                'ok': False,
                'error': 'Unauthorized: Invalid or missing cron secret.'
            }, status=401)

    # Optional query parameters: force, test, date (YYYY-MM-DD)
    force = request.GET.get('force') == 'true' or request.POST.get('force') == 'true'
    is_test = request.GET.get('test') == 'true' or request.POST.get('test') == 'true'
    date_str = request.GET.get('date') or request.POST.get('date')
    target_date = None
    if date_str:
        try:
            target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return JsonResponse({
                'ok': False,
                'error': f"Invalid date format '{date_str}'. Expected YYYY-MM-DD."
            }, status=400)

    base_url = request.build_absolute_uri('/')
    result = send_daily_summary_email(
        target_date=target_date,
        force=force,
        is_test=is_test,
        base_url=base_url
    )

    status_code = 200 if result.get('ok') else 400
    if result.get('reason') in ('disabled', 'zero_feedback_suppressed'):
        status_code = 200

    return JsonResponse(result, status=status_code)


@superuser_required
@require_POST
def backup_create(request):
    if not request.user.is_superuser:
        return JsonResponse({'ok': False, 'error': 'Only superusers can create backups.'}, status=403)
    try:
        result = create_backup()
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Backup failed: {e}'}, status=500)

    log_admin_event(request.user, FeedbackConfiguration.get_solo(), ADDITION,
                     f'Created backup "{result["filename"]}"')

    feedback_count = FeedbackEntry.objects.count()

    return JsonResponse({
        'ok': True,
        'message': f'Backup created: {result["filename"]}',
        'feedback_count': feedback_count,
        'backup': {
            'filename': result['filename'],
            'size_bytes': result['size_bytes'],
            'size_display': result.get('size_display', ''),
            'created_display': result.get('created_display', ''),
            'download_url': reverse('backup_download', args=[result['filename']]),
        },
    })


@superuser_required
def backup_download(request, filename):
    if not request.user.is_superuser:
        messages.error(request, 'Only superusers can download backups.')
        return redirect('settings_page')
    try:
        path = resolve_backup_path(filename)
    except (SuspiciousFileOperation, FileNotFoundError):
        messages.error(request, f'Backup file "{filename}" was not found on disk. The list has been refreshed.')
        return redirect(reverse('settings_page') + '#system-settings')
    return FileResponse(open(path, 'rb'), as_attachment=True, filename=filename)


@superuser_required
@require_POST
def backup_delete(request, filename):
    if not request.user.is_superuser:
        return JsonResponse({'ok': False, 'error': 'Only superusers can delete backups.'}, status=403)
    try:
        delete_backup(filename)
    except (SuspiciousFileOperation, FileNotFoundError):
        return JsonResponse({'ok': False, 'error': 'Backup not found.'}, status=404)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Delete failed: {e}'}, status=500)

    log_admin_event(request.user, FeedbackConfiguration.get_solo(), DELETION,
                     f'Deleted backup "{filename}"')
    return JsonResponse({'ok': True, 'message': f'Backup "{filename}" deleted.'})


@superuser_required
@require_POST
def backup_restore(request):
    if not request.user.is_superuser:
        return JsonResponse({'ok': False, 'error': 'Only superusers can restore backups.'}, status=403)

    uploaded = request.FILES.get('backup_file')
    existing_filename = request.POST.get('existing_filename', '').strip()

    if not uploaded and not existing_filename:
        return JsonResponse({'ok': False, 'error': 'No backup file or archive specified.'}, status=400)

    try:
        if uploaded:
            safety = restore_backup(uploaded)
            source_label = getattr(uploaded, 'name', 'uploaded backup')
        else:
            path = resolve_backup_path(existing_filename)
            with open(path, 'rb') as f:
                safety = restore_backup(f)
            source_label = existing_filename
    except (SuspiciousFileOperation, FileNotFoundError):
        return JsonResponse({'ok': False, 'error': f'Backup file "{existing_filename}" was not found on disk.'}, status=404)
    except ValueError as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Restore failed: {e}'}, status=500)

    log_admin_event(request.user, FeedbackConfiguration.get_solo(), CHANGE,
                     f'Restored database from "{source_label}" (safety backup: "{safety["filename"]}")')

    feedback_count = FeedbackEntry.objects.count()

    return JsonResponse({
        'ok': True,
        'message': f'Database restored successfully from "{source_label}". Previous data was saved as "{safety["filename"]}".',
        'feedback_count': feedback_count,
        'safety_backup': {
            'filename': safety['filename'],
            'size_bytes': safety['size_bytes'],
            'size_display': safety.get('size_display', ''),
            'created_display': safety.get('created_display', ''),
            'download_url': reverse('backup_download', args=[safety['filename']]),
        },
    })